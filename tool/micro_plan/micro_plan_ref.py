"""Reference implementation of the Micro Plan parsing rules.

Runs against the real workbook so every rule is checked on real data before
it is ported to Dart. Output: per-sheet counts that the Dart tests must match.
"""
import datetime
import json
import re
import sys
from collections import Counter

import openpyxl

MONTHS = {
    'APRIL': 4, 'MAY': 5, 'JUNE': 6, 'JULY': 7, 'AUG': 8, 'SEP': 9,
    'OCT': 10, 'NOV': 11, 'DEC': 12, 'JAN': 1, 'FEB': 2, 'MARCH': 3,
}
DAYS = ['MONDAY', 'TUESDAY', 'WEDNESDAY', 'THURSDAY', 'FRIDAY', 'SATURDAY',
        'SUNDAY']
HEADER_ROW = 12
FIRST_DATA_ROW = 14
LAST_COL = 14  # A..N
VERTICAL_FILL_COLS = [1, 2, 13, 14]  # S.No, Name, Visit date, Day


def sheet_month(name):
    m = re.fullmatch(r'\s*([A-Z]+)\s*(\d{2})\s*', name.upper())
    if not m or m.group(1) not in MONTHS:
        return None
    return MONTHS[m.group(1)], 2000 + int(m.group(2))


def text(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def to_int(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return int(v) if v.is_integer() else None
    s = str(v).strip()
    return int(s) if re.fullmatch(r'-?\d+', s) else None


def resolve_date(raw, month, year, day_name, flags):
    """Return a date that lies in the sheet's month, or None + a flag."""
    cands = []
    if isinstance(raw, datetime.datetime):
        cands.append((raw.year, raw.month, raw.day))
        cands.append((raw.year, raw.day, raw.month))  # Excel day/month swap
    elif isinstance(raw, str):
        if not raw.strip():
            return None
        m = re.fullmatch(r'\s*(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{4})\s*', raw)
        if not m:
            flags.append('DATE_UNREADABLE')
            return None
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        cands.append((y, mo, d))
    elif raw is not None:
        if text(raw) is None:
            return None
        flags.append('DATE_UNREADABLE')
        return None
    else:
        return None
    valid = []
    for y, mo, d in cands:
        try:
            valid.append(datetime.date(y, mo, d))
        except ValueError:
            pass
    in_month = [c for c in valid if c.month == month and c.year == year]
    in_month = list(dict.fromkeys(in_month))
    if not in_month:
        flags.append('DATE_OUTSIDE_SHEET_MONTH')
        return None
    if len(in_month) > 1:  # only possible when day == month, same date anyway
        in_month = in_month[:1]
    date = in_month[0]
    if isinstance(raw, datetime.datetime) and (raw.month, raw.day) != (date.month, date.day):
        flags.append('DATE_DAY_MONTH_SWAPPED')
    if day_name and day_name.upper() in DAYS and DAYS[date.weekday()] != day_name.upper():
        flags.append('DAY_NAME_MISMATCH')
    return date


def parse_sheet(ws):
    sm = sheet_month(ws.title)
    grid = {}
    for r in range(1, ws.max_row + 1 if ws.max_row < 400 else 400):
        for c in range(1, LAST_COL + 1):
            grid[(r, c)] = ws.cell(r, c).value
    # vertical merges: copy the top cell's value into the cells below it
    filled = set()
    for m in ws.merged_cells.ranges:
        top = grid.get((m.min_row, m.min_col))
        for r in range(m.min_row + 1, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                if c in VERTICAL_FILL_COLS and m.min_col == m.max_col and r >= FIRST_DATA_ROW:
                    grid[(r, c)] = top
                    filled.add((r, c))
    rows = []
    for r in range(FIRST_DATA_ROW, 400):
        vals = [grid.get((r, c)) for c in range(1, LAST_COL + 1)]
        if all(text(v) is None for v in vals):
            continue
        flags = []
        sno, name, typ = to_int(vals[0]), text(vals[1]), text(vals[2])
        typ_u = typ.upper() if typ else None
        name_u = name.upper() if name else ''
        if typ_u == 'SCHOOL':
            kind = 'SCHOOL'
        elif typ_u == 'AWC':
            kind = 'AWC'
        elif typ_u is not None:
            kind = 'OTHER'
            flags.append('UNKNOWN_INSTITUTION_TYPE')
        elif name_u == 'SUNDAY':
            kind = 'SUNDAY'
        elif name_u == 'PHC REFERRED CHILDREN TREATMENT':
            kind = 'EVENT'
        elif name:
            kind = 'HOLIDAY'
        else:
            kind = 'OTHER'
            flags.append('UNRECOGNISED_ROW')
        date = None
        if sm:
            date = resolve_date(vals[12], sm[0], sm[1], text(vals[13]), flags)
        if date is None and 'DATE_UNREADABLE' not in flags and 'DATE_OUTSIDE_SHEET_MONTH' not in flags and kind != 'OTHER':
            flags.append('DATE_MISSING')
        if (r, 2) in filled and kind in ('SCHOOL', 'AWC'):
            flags.append('NAME_SHARED_WITH_ROW_ABOVE')
        male, female, total = to_int(vals[7]), to_int(vals[8]), to_int(vals[9])
        if kind in ('SCHOOL', 'AWC'):
            if not name:
                flags.append('NAME_MISSING')
            if male is None or female is None or total is None:
                flags.append('COUNT_MISSING')
            elif male + female != total:
                flags.append('COUNT_TOTAL_MISMATCH')
            phone = text(vals[11])
            if phone and not re.fullmatch(r'\d{10}', phone):
                flags.append('CONTACT_NUMBER_NOT_10_DIGITS')
            if kind == 'SCHOOL':
                code = text(vals[4])
                if code is None:
                    flags.append('SCHOOL_CODE_MISSING')
                elif not re.fullmatch(r'\d{10}', code):
                    flags.append('SCHOOL_CODE_UNUSUAL')
            if kind == 'AWC' and text(vals[3]) is None:
                flags.append('AWC_CODE_MISSING')
        rows.append({'row': r, 'kind': kind, 'sno': sno, 'name': name,
                     'date': date.isoformat() if date else None,
                     'flags': flags})
    return rows


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    summary = {}
    all_flags = Counter()
    for ws in wb.worksheets:
        rows = parse_sheet(ws)
        kinds = Counter(r['kind'] for r in rows)
        fl = Counter(f for r in rows for f in r['flags'])
        all_flags.update(fl)
        summary[ws.title] = {'rows': len(rows), 'kinds': dict(kinds),
                             'flags': dict(fl)}
    print(json.dumps(summary, indent=1))
    print('ALL FLAGS', all_flags)
    return wb


if __name__ == '__main__':
    main(sys.argv[1])
