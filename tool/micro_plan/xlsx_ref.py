"""Raw .xlsx reader — the exact steps the Dart XlsxWorkbookReader takes.

Checked against openpyxl on the real workbook: every cell value and every
merged range must agree. Usage: python xlsx_ref.py <file.xlsx>
"""
import datetime
import re
import sys
import zipfile
import xml.etree.ElementTree as ET

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
REL_NS = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
PKG_REL = '{http://schemas.openxmlformats.org/package/2006/relationships}'
BUILTIN_DATE_FORMATS = set(range(14, 23)) | {45, 46, 47}


def col_index(letters):
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def split_ref(ref):
    m = re.fullmatch(r'([A-Z]+)(\d+)', ref)
    return int(m.group(2)), col_index(m.group(1))  # 1-based row, 0-based col


def is_date_code(code):
    code = re.sub(r'"[^"]*"', '', code)
    code = re.sub(r'\[[^\]]*\]', '', code)
    code = re.sub(r'\\.', '', code)
    return re.search(r'[dmyhs]', code, re.I) is not None


def text_of(si):
    # Plain <t>, or the <t> of every <r> run; phonetic <rPh> is skipped.
    parts = []
    for child in si:
        if child.tag == NS + 't':
            parts.append(child.text or '')
        elif child.tag == NS + 'r':
            t = child.find(NS + 't')
            parts.append(t.text or '' if t is not None else '')
    return ''.join(parts)


def serial_to_date(serial, date1904):
    if date1904:
        base = datetime.datetime(1904, 1, 1)
    elif serial < 60:
        base = datetime.datetime(1899, 12, 31)
    else:
        base = datetime.datetime(1899, 12, 30)
    ms = round(serial * 86400000)
    return base + datetime.timedelta(milliseconds=ms)


def number(text):
    if re.fullmatch(r'-?\d+', text):
        return int(text)
    return float(text)


def read(path):
    z = zipfile.ZipFile(path)
    wb = ET.fromstring(z.read('xl/workbook.xml'))
    pr = wb.find(NS + 'workbookPr')
    date1904 = pr is not None and pr.get('date1904') in ('1', 'true')
    rels = ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
    targets = {r.get('Id'): r.get('Target') for r in rels.iter(PKG_REL + 'Relationship')}

    shared = []
    if 'xl/sharedStrings.xml' in z.namelist():
        sst = ET.fromstring(z.read('xl/sharedStrings.xml'))
        shared = [text_of(si) for si in sst.findall(NS + 'si')]

    date_styles = set()
    if 'xl/styles.xml' in z.namelist():
        st = ET.fromstring(z.read('xl/styles.xml'))
        custom = {}
        nf = st.find(NS + 'numFmts')
        if nf is not None:
            for f in nf.findall(NS + 'numFmt'):
                custom[int(f.get('numFmtId'))] = f.get('formatCode') or ''
        xfs = st.find(NS + 'cellXfs')
        if xfs is not None:
            for i, xf in enumerate(xfs.findall(NS + 'xf')):
                fid = int(xf.get('numFmtId') or 0)
                if fid in BUILTIN_DATE_FORMATS or (fid in custom and is_date_code(custom[fid])):
                    date_styles.add(i)

    sheets = []
    for sh in wb.find(NS + 'sheets').findall(NS + 'sheet'):
        target = targets[sh.get(REL_NS + 'id')].lstrip('/')
        if not target.startswith('xl/'):
            target = 'xl/' + target
        root = ET.fromstring(z.read(target))
        rows = {}
        next_row = 1
        for row in root.iter(NS + 'row'):
            r = int(row.get('r')) if row.get('r') else next_row
            next_row = r + 1
            cells = rows.setdefault(r, {})
            next_col = 0
            for c in row.findall(NS + 'c'):
                if c.get('r'):
                    _, col = split_ref(c.get('r'))
                else:
                    col = next_col
                next_col = col + 1
                t = c.get('t')
                v = c.find(NS + 'v')
                vt = v.text if v is not None else None
                if t == 's':
                    value = shared[int(vt)] if vt is not None else None
                elif t == 'inlineStr':
                    is_ = c.find(NS + 'is')
                    value = text_of(is_) if is_ is not None else None
                elif t in ('str', 'e'):
                    value = vt
                elif t == 'b':
                    value = None if vt is None else vt == '1'
                elif t == 'd':
                    value = datetime.datetime.fromisoformat(vt) if vt else None
                else:
                    if vt is None or vt == '':
                        value = None
                    else:
                        value = number(vt)
                        if int(c.get('s') or 0) in date_styles:
                            value = serial_to_date(float(vt), date1904)
                cells[col] = value
        merges = []
        mc = root.find(NS + 'mergeCells')
        if mc is not None:
            for m in mc.findall(NS + 'mergeCell'):
                a, b = m.get('ref').split(':')
                r1, c1 = split_ref(a)
                r2, c2 = split_ref(b)
                merges.append((r1, r2, c1 + 1, c2 + 1))
        # Only the top-left cell of a merged range has a value (Excel keeps
        # stale cached values in the covered cells; they are not shown).
        for r1, r2, c1, c2 in merges:
            for r in range(r1, r2 + 1):
                for c in range(c1 - 1, c2):
                    if (r, c) != (r1, c1 - 1) and c in rows.get(r, {}):
                        rows[r][c] = None
        # Trailing empty rows (formatted but blank) are dropped.
        filled = [r for r, cells in rows.items()
                  if any(v is not None and v != '' for v in cells.values())]
        max_row = max(filled) if filled else 0
        grid = []
        for r in range(1, max_row + 1):
            cells = rows.get(r, {})
            used = [c for c, v in cells.items() if v is not None]
            width = max(used) + 1 if used else 0
            grid.append([cells.get(c) for c in range(width)])
        sheets.append((sh.get('name'), grid, merges))
    return sheets


def compare_with_openpyxl(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ours = read(path)
    problems = 0
    for (name, grid, merges), ws in zip(ours, wb.worksheets):
        assert name == ws.title, (name, ws.title)
        # Compare the area either side has content in (sheets report huge
        # max_column values when whole columns are formatted).
        last_row = max(len(grid), min(ws.max_row, 400))
        for r in range(1, last_row + 1):
            for c in range(1, 27):
                theirs = ws.cell(r, c).value
                row = grid[r - 1] if r - 1 < len(grid) else []
                mine = row[c - 1] if c - 1 < len(row) else None
                if theirs != mine or type(theirs) != type(mine):
                    problems += 1
                    if problems < 15:
                        print('DIFF', name, ws.cell(r, c).coordinate, repr(theirs), repr(mine))
        theirs_m = sorted((m.min_row, m.max_row, m.min_col, m.max_col) for m in ws.merged_cells.ranges)
        if sorted(merges) != theirs_m:
            problems += 1
            print('MERGE DIFF', name)
    print('sheets', len(ours), 'problems', problems)
    return problems


if __name__ == '__main__':
    sys.exit(1 if compare_with_openpyxl(sys.argv[1]) else 0)
